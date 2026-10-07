"""Importa lo scadenziario dei contratti a termine (Excel dello studio) nella collezione `contratti`.
Uso: python3 -I gen/contratti_termine.py <file.xlsx> <cartella clienti json> <cartella uscita>"""
import openpyxl, datetime, json, re, sys, os, hashlib, glob

XL, CLD, OUT = sys.argv[1:4]
os.makedirs(OUT, exist_ok=True)
CL = {os.path.basename(f)[:-5]: (lambda x: x.get('data', x))(json.load(open(f))) for f in glob.glob(f'{CLD}/*.json')}
FONTE = "Scadenziario Excel aggiornato al 30/09/2026"
FLAGS = {"STAGIONALE": "stagionale", "CAUSALE": "causale", "INTERMITTENTE": "intermittente"}

def pdate(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})\s*", v)
        if m:
            d, mo, y = map(int, m.groups()); y += 2000 if y < 100 else 0
            return datetime.date(y, mo, d).isoformat()
    return None

def split_az(a):
    a = a.strip(); sede, note = "", ""
    m = re.search(r"_?\s*\(([^)]*)\)\s*$", a)
    if m:
        inner = m.group(1).strip(); a = a[:m.start()].strip(" _")
        if inner.upper().startswith("PROR"): note = inner
        else: sede = inner
    elif "_" in a:
        a, sede = [x.strip() for x in a.split("_", 1)]
    return a, sede, note

def split_dip(d):
    d = d.strip(); tipo = "td"
    for pat, t in [(r"[_\s]*\(?co\.co\.co\.?\)?$", "cococo"), (r"[_\s]*\(?tiroc\.?\)?$", "tirocinio"), (r"[_\s]*distacco$", "distacco")]:
        m = re.search(pat, d, re.I)
        if m: d, tipo = d[:m.start()].strip(" _"), t; break
    return d, tipo

docs = []
def add(cliente, azienda, sede, dip, tipo, ass, cells, note_extra=""):
    scad, notes, fl = [], [], set()
    for v in cells:
        dt = pdate(v)
        if dt: scad.append(dt); continue
        if v is None or (isinstance(v, str) and not v.strip()): continue
        s = str(v).strip()
        if s == "-": notes.append("Fine iniziale non indicata"); continue
        if s.upper() in FLAGS: fl.add(FLAGS[s.upper()]); continue
        notes.append(s)
    if note_extra: notes.insert(0, note_extra)
    if "intermittente" in fl and tipo == "td": tipo = "intermittente"
    d = {"cliente": cliente, "azienda": azienda, "sede": sede, "dipendente": dip, "tipo": tipo,
         "stagionale": "stagionale" in fl, "causale": "causale" in fl, "assunzione": pdate(ass) or "",
         "scadenze": scad, "note": " · ".join(notes), "stato": "attivo", "fonte": FONTE,
         "aggiornatoIl": "2026-10-07T12:00:00+02:00", "aggiornatoDa": ""}
    key = f"{cliente}|{azienda}|{sede}|{dip}|{d['assunzione']}"
    d_id = "t-" + hashlib.sha1(key.encode()).hexdigest()[:12]
    docs.append((d_id, d))

wb = openpyxl.load_workbook(XL, data_only=True)
ws = wb["Contratti"]
for r in ws.iter_rows(min_row=6, values_only=True):
    n, az, dip, ass = r[0], r[1], r[2], r[3]
    if not az or not dip: continue
    cliente = "c-leucopetra" if n == "*" else f"c-{n}"
    a, sede, note = split_az(az); name, tipo = split_dip(dip)
    add(cliente, a, sede, name, tipo, ass, r[4:], note)

ws = wb["Beducci Travel Bus Srl"]
for r in ws.iter_rows(min_row=5, values_only=True):
    if not r[0] or not r[2]: continue
    fil, cog, nom = (r[1] or "").strip(), (r[2] or "").strip(), (r[3] or "").strip()
    if fil.upper() in ("AUTORIMESSE", "TURISMO", "AUTOFERROTRANVIERI"): sede, dip = fil.capitalize(), f"{cog.title()} {nom.title()}"
    else: sede, dip = "", f"{fil.title()} {cog.title()} {nom.title()}"
    add("c-344", "BEDUCCI TRAVEL BUS SRL", sede, dip.strip(), "td", r[4], list(r[5:10]) + [r[10]])

ws = wb[[s for s in wb.sheetnames if s.startswith("Terrazza")][0]]
for r in ws.iter_rows(min_row=5, values_only=True):
    if not r[0] or not r[1]: continue
    add("c-345", "TERRAZZA DUE GOLFI SRL", "", f"{r[1].strip().title()} {r[2].strip().title()}", "td", r[3], list(r[4:10]))

ids = [i for i, _ in docs]; assert len(ids) == len(set(ids)), "id duplicati"
for i, d in docs: json.dump(d, open(f"{OUT}/{i}.json", "w"), ensure_ascii=False)
w = [{"op": "set", "collection": "contratti", "doc_id": i, "file_path": os.path.abspath(f"{OUT}/{i}.json")} for i, _ in docs]
json.dump([w[k:k+50] for k in range(0, len(w), 50)], open(f"{OUT}/batches.json", "w"))
print(len(docs), "contratti")
miss = sorted({d["cliente"] for _, d in docs if d["cliente"] not in CL}); print("clienti non in anagrafica:", [(c, next(d["azienda"] for _, d in docs if d["cliente"] == c)) for c in miss])
import collections
print(collections.Counter(d["tipo"] for _, d in docs), "stagionali", sum(d["stagionale"] for _, d in docs))
for _, d in docs:
    if not d["scadenze"] or d["scadenze"] != sorted(d["scadenze"]) or (d["assunzione"] and d["scadenze"] and d["scadenze"][0] < d["assunzione"]):
        print("DATE ANOMALE", d["azienda"], d["dipendente"], d["assunzione"], d["scadenze"])
print("note:", sorted({d["note"] for _, d in docs if d["note"]}))
print("sedi:", sorted({(d["azienda"], d["sede"]) for _, d in docs})[:80])
